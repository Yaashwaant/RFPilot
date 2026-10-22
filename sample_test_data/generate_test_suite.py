#!/usr/bin/env python3
"""
Generates comprehensive sample test documents and RFP Excel workbooks for testing RFPilot:
1. ApexFlow_Security_and_Architecture_Whitepaper.pdf (Multi-page PDF for Knowledge upload)
2. ApexFlow_Security_and_Architecture_Whitepaper.txt (Text file for Knowledge upload)
3. ApexFlow_Enterprise_Trust_Center.md (Markdown document for Knowledge upload)
4. RFP_Evaluation_Worksheet_Blank.xlsx (The RFP questionnaire to upload to a Project and have RFPilot fill out)
5. Past_Completed_RFP_ApexFlow_2025.xlsx (Past completed RFP with answers to test Connector / Mem0 ingestion)
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

WHITE_PAPER_TEXT = """ApexFlow Technologies – Enterprise Security & Architecture Whitepaper (2026)

1. Executive Summary & Architecture Overview
ApexFlow Enterprise Suite v4.2 is a cloud-native B2B workflow automation platform designed for global enterprises. The architecture is deployed across multi-region high-availability clusters in AWS US-East (Northern Virginia) and AWS EU-Central (Frankfurt). The platform operates an active-active failover model with automated multi-zone replication, delivering high resiliency and zero single points of failure.

2. Authentication, Single Sign-On (SSO) & Identity Management
ApexFlow supports turnkey enterprise Single Sign-On (SSO) via SAML 2.0 and OpenID Connect (OIDC). Verified turnkey identity providers include Okta, Microsoft Entra ID (Azure AD), Google Workspace, PingFederate, and CyberArk.
- User Provisioning: Just-In-Time (JIT) provisioning and SCIM 2.0 directory synchronization for automated user on-boarding, role mapping, and instantaneous de-provisioning upon employee offboarding.
- Multi-Factor Authentication: MFA is enforced by default, supporting FIDO2 / WebAuthn, hardware security keys (YubiKey), and TOTP authenticator apps.
- Role-Based Access Control: Granular RBAC supporting custom enterprise permission tiers (Global Admin, Department Head, Approver, Analyst, Auditor, Read-Only).

3. Data Encryption, Cryptography & Tenant Isolation
All enterprise customer data is strictly segregated using isolated logical tenant schemas, database-level Row-Level Security (RLS), and dedicated cryptographic tenant partition keys.
- Encryption at Rest: All production databases, cache layers, and object storage volumes are encrypted using AES-256 with Customer-Managed Keys (BYOK) via AWS KMS. Master encryption keys are rotated automatically every 365 days.
- Encryption in Transit: TLS 1.3 is enforced with HTTP Strict Transport Security (HSTS) on all external web, mobile, and API endpoints. Internal service-to-service microservice communications require mutual TLS (mTLS) with SPIFFE identity certificates.

4. Data Retention, Backup & Disposal Standards
Customer operational data is retained for 90 days following subscription termination by default, after which a certified cryptographic wipe conforming to NIST SP 800-88 Rev 1 is executed with an auditable certificate of destruction provided upon request.
- Automated Backups: Database snapshots are captured every 24 hours with continuous write-ahead logging (WAL), retained for 35 days, encrypted with AES-256.
- Cross-Region Replication: Daily encrypted backups are replicated to an off-site secondary disaster recovery region.

5. Compliance Certifications, Audits & Regulatory Adherence
ApexFlow maintains rigorous third-party audited compliance certifications:
- SOC 2 Type II: Annual audit conducted by Ernst & Young covering Security, Availability, and Confidentiality trust service criteria. Latest report issued November 2025.
- ISO Certifications: Certified under ISO/IEC 27001:2022 (Information Security Management) and ISO/IEC 27701:2019 (Privacy Information Management).
- Privacy Regulations: Fully compliant with EU General Data Protection Regulation (GDPR) and California Consumer Privacy Act (CCPA/CPRA). Standard Contractual Clauses (SCCs) and Data Processing Agreements (DPA) are executed for all enterprise engagements.
- Healthcare Compliance: Compliant with HIPAA Security & Privacy Rules; ApexFlow signs standard Business Associate Agreements (BAAs) for covered entities.

6. Service Level Agreement (SLA), Availability & Disaster Recovery
- Monthly Uptime Guarantee: ApexFlow guarantees 99.95% monthly application uptime under our Enterprise SLA, backed by tiered financial service credits for any downtime.
- Recovery Time Objective (RTO): Guaranteed RTO of less than 1 hour in the event of an infrastructure or regional disaster.
- Recovery Point Objective (RPO): Guaranteed RPO of less than 15 minutes across all customer transactional data stores.
- DR Exercises: Disaster recovery simulations and table-top exercises are performed bi-annually with documented executive post-mortems.

7. Enterprise APIs, Webhooks & Rate Limiting
- REST & GraphQL: ApexFlow provides a complete RESTful API compliant with OpenAPI 3.1 specifications as well as a GraphQL query interface.
- Rate Limits: Standard enterprise API quotas allow up to 10,000 requests per minute per tenant, governed by a token-bucket algorithm with standard HTTP 429 Retry-After response headers.
- Webhooks & Events: Outbound webhook notifications with HMAC-SHA256 signature verification allow real-time event streaming to enterprise Slack, Jira, ServiceNow, and custom webhook endpoints.
"""

def generate_pdf():
    pdf_path = os.path.join(OUT_DIR, "ApexFlow_Security_and_Architecture_Whitepaper.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=15,
    )
    h2_style = ParagraphStyle(
        'SectionH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#1E3A8A'),
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6,
    )
    bullet_style = ParagraphStyle(
        'BulletPoint',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#334155'),
        leftIndent=15,
        spaceAfter=4,
    )

    story = []
    story.append(Paragraph("ApexFlow Technologies – Enterprise Security & Architecture", title_style))
    story.append(Paragraph("Document Version: 4.2 | Classification: Confidential / External Client Review | Date: January 2026", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#3B82F6'), spaceAfter=14))

    sections = [
        ("1. Architecture & Multi-Region Hosting", [
            "ApexFlow Enterprise Suite v4.2 is a cloud-native B2B workflow automation platform designed for global enterprises. The architecture is deployed across multi-region high-availability clusters in AWS US-East (Northern Virginia) and AWS EU-Central (Frankfurt).",
            "The platform operates an active-active failover model with automated multi-zone replication, delivering high resiliency and zero single points of failure.",
        ]),
        ("2. Authentication, Single Sign-On (SSO) & SCIM", [
            "ApexFlow supports turnkey enterprise Single Sign-On (SSO) via SAML 2.0 and OpenID Connect (OIDC). Verified turnkey identity providers include Okta, Microsoft Entra ID (Azure AD), Google Workspace, PingFederate, and CyberArk.",
            "• <b>User Provisioning:</b> Just-In-Time (JIT) provisioning and SCIM 2.0 directory synchronization for automated user on-boarding, role mapping, and instantaneous de-provisioning upon employee offboarding.",
            "• <b>Multi-Factor Authentication (MFA):</b> MFA is enforced by default, supporting FIDO2 / WebAuthn, hardware security keys (YubiKey), and TOTP authenticator apps.",
            "• <b>Role-Based Access Control:</b> Granular RBAC supporting custom enterprise permission tiers (Global Admin, Department Head, Approver, Analyst, Auditor, Read-Only).",
        ]),
        ("3. Data Encryption & Tenant Isolation", [
            "All enterprise customer data is strictly segregated using isolated logical tenant schemas, database-level Row-Level Security (RLS), and dedicated cryptographic tenant partition keys.",
            "• <b>Encryption at Rest:</b> All production databases, cache layers, and object storage volumes are encrypted using AES-256 with Customer-Managed Keys (BYOK) via AWS KMS. Master encryption keys are rotated automatically every 365 days.",
            "• <b>Encryption in Transit:</b> TLS 1.3 is enforced with HTTP Strict Transport Security (HSTS) on all external web, mobile, and API endpoints. Internal service-to-service microservice communications require mutual TLS (mTLS) with SPIFFE identity certificates.",
        ]),
        ("4. Data Retention, Backup & Disposal", [
            "Customer operational data is retained for 90 days following subscription termination by default, after which a certified cryptographic wipe conforming to NIST SP 800-88 Rev 1 is executed with an auditable certificate of destruction provided upon request.",
            "• <b>Automated Backups:</b> Database snapshots are captured every 24 hours with continuous write-ahead logging (WAL), retained for 35 days, encrypted with AES-256.",
            "• <b>Cross-Region Replication:</b> Daily encrypted backups are replicated to an off-site secondary disaster recovery region.",
        ]),
        ("5. Compliance Certifications & Audits", [
            "ApexFlow maintains rigorous third-party audited compliance certifications:",
            "• <b>SOC 2 Type II:</b> Annual audit conducted by Ernst & Young covering Security, Availability, and Confidentiality trust service criteria. Latest report issued November 2025.",
            "• <b>ISO Certifications:</b> Certified under ISO/IEC 27001:2022 (Information Security Management) and ISO/IEC 27701:2019 (Privacy Information Management).",
            "• <b>Privacy Regulations:</b> Fully compliant with EU General Data Protection Regulation (GDPR) and California Consumer Privacy Act (CCPA/CPRA). Standard Contractual Clauses (SCCs) and Data Processing Agreements (DPA) are executed for all enterprise engagements.",
            "• <b>Healthcare Compliance:</b> Compliant with HIPAA Security & Privacy Rules; ApexFlow signs standard Business Associate Agreements (BAAs) for covered entities.",
        ]),
        ("6. SLAs, Availability & Disaster Recovery", [
            "• <b>Monthly Uptime Guarantee:</b> ApexFlow guarantees 99.95% monthly application uptime under our Enterprise SLA, backed by tiered financial service credits for any downtime.",
            "• <b>Recovery Time Objective (RTO):</b> Guaranteed RTO of less than 1 hour in the event of an infrastructure or regional disaster.",
            "• <b>Recovery Point Objective (RPO):</b> Guaranteed RPO of less than 15 minutes across all customer transactional data stores.",
            "• <b>DR Simulations:</b> Disaster recovery simulations and table-top exercises are performed bi-annually with documented executive post-mortems.",
        ]),
        ("7. Enterprise APIs, Webhooks & Rate Limiting", [
            "• <b>REST & GraphQL:</b> ApexFlow provides a complete RESTful API compliant with OpenAPI 3.1 specifications as well as a GraphQL query interface.",
            "• <b>Rate Limits:</b> Standard enterprise API quotas allow up to 10,000 requests per minute per tenant, governed by a token-bucket algorithm with standard HTTP 429 Retry-After response headers.",
            "• <b>Webhooks & Events:</b> Outbound webhook notifications with HMAC-SHA256 signature verification allow real-time event streaming to enterprise Slack, Jira, ServiceNow, and custom webhook endpoints.",
        ]),
    ]

    for title, paras in sections:
        story.append(Paragraph(title, h2_style))
        for p in paras:
            if p.startswith("•"):
                story.append(Paragraph(p, bullet_style))
            else:
                story.append(Paragraph(p, body_style))
        story.append(Spacer(1, 4))

    doc.build(story)
    print(f"Created: {pdf_path}")

def generate_text_and_md():
    txt_path = os.path.join(OUT_DIR, "ApexFlow_Security_and_Architecture_Whitepaper.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(WHITE_PAPER_TEXT)
    print(f"Created: {txt_path}")

    md_path = os.path.join(OUT_DIR, "ApexFlow_Enterprise_Trust_Center.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(f"# ApexFlow Enterprise Trust Center & Architecture Specifications\n\n{WHITE_PAPER_TEXT}")
    print(f"Created: {md_path}")

def generate_blank_rfp_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Security & Architecture RFP"

    # Styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    border_side = Side(style='thin', color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    headers = [
        "Requirement ID",
        "Category",
        "RFP Question / Technical Requirement",
        "Vendor Response (To be filled by Proposal Assistant)",
        "Compliance Status",
        "Notes / Cross-Reference",
    ]
    ws.append(headers)

    questions = [
        ("REQ-01", "Authentication & SSO", "Does your application support Single Sign-On (SSO) via SAML 2.0 or OIDC with enterprise identity providers like Okta and Microsoft Entra ID?", "", "Compliant", ""),
        ("REQ-02", "User Lifecycle", "Do you support automated user provisioning and de-provisioning via SCIM 2.0 or Just-In-Time (JIT) provisioning?", "", "Compliant", ""),
        ("REQ-03", "Data Encryption", "How is customer data encrypted both at rest and in transit? Please specify cryptographic algorithms and key management practices.", "", "Compliant", ""),
        ("REQ-04", "Data Retention & Disposal", "What is your customer data retention policy following contract termination, and what certified standards do you follow for data sanitization?", "", "Compliant", ""),
        ("REQ-05", "Compliance & Certifications", "Do you maintain an active SOC 2 Type II report and ISO/IEC 27001 certification? Who is your third-party auditor and when was the last report issued?", "", "Compliant", ""),
        ("REQ-06", "Data Privacy & Governance", "Are you compliant with EU GDPR and California CCPA/CPRA? Do you execute standard Data Processing Agreements (DPA) and Standard Contractual Clauses (SCCs)?", "", "Compliant", ""),
        ("REQ-07", "Healthcare Compliance", "Can your solution support healthcare data in accordance with HIPAA, and will you sign a standard Business Associate Agreement (BAA)?", "", "Compliant", ""),
        ("REQ-08", "Service Level Agreement (SLA)", "What is your guaranteed monthly uptime SLA, and do you offer financial service credits in the event of an outage?", "", "Compliant", ""),
        ("REQ-09", "Disaster Recovery", "What are your guaranteed Recovery Time Objective (RTO) and Recovery Point Objective (RPO) metrics for disaster recovery?", "", "Compliant", ""),
        ("REQ-10", "Tenant Architecture", "Describe how multi-tenant data isolation is enforced at the database and application tiers to prevent cross-tenant data leakage.", "", "Compliant", ""),
        ("REQ-11", "API & Webhook Integrations", "Does your platform offer REST or GraphQL APIs, and what are the standard enterprise rate limits and webhook notification capabilities?", "", "Compliant", ""),
        ("REQ-12", "Backup & Recovery", "How frequently are database backups performed, how long are they retained, and are they stored offsite in an encrypted state?", "", "Compliant", ""),
    ]

    for row_idx, q in enumerate(questions, start=2):
        ws.append(q)

    # Style header row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx in (1, 5) else "left", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws.row_dimensions[1].height = 28

    # Style data rows
    for row_idx in range(2, len(questions) + 2):
        ws.row_dimensions[row_idx].height = 40
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = bold_font if col_idx in (1, 2) else data_font
            cell.border = cell_border
            if is_even and col_idx != 4:
                cell.fill = alt_fill
            # Alignments
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (3, 4):
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column widths
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 46
    ws.column_dimensions['D'].width = 54  # Vendor response column
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 22

    xlsx_path = os.path.join(OUT_DIR, "RFP_Evaluation_Worksheet_Blank.xlsx")
    wb.save(xlsx_path)
    print(f"Created: {xlsx_path}")

def generate_past_rfp_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Past Approved RFP Responses"

    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    border_side = Side(style='thin', color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    alt_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")

    headers = [
        "Req ID",
        "Category",
        "Requirement / Question",
        "Approved Vendor Answer",
        "Approval Status",
    ]
    ws.append(headers)

    past_qa = [
        ("SEC-01", "Authentication", "Describe your Single Sign-On and multi-factor authentication support.", "ApexFlow supports turnkey enterprise Single Sign-On (SSO) via SAML 2.0 and OpenID Connect (OIDC), pre-integrated with Okta, Microsoft Entra ID (Azure AD), Google Workspace, and PingFederate. Multi-factor authentication (MFA) is strictly enforced with support for FIDO2/WebAuthn hardware keys and TOTP authenticators.", "APPROVED"),
        ("SEC-02", "Encryption", "What encryption mechanisms protect customer data?", "All customer data at rest is encrypted using AES-256 with Customer-Managed Keys (BYOK) via AWS KMS. Data in transit is secured with TLS 1.3 across public endpoints, and internal microservice communication is protected via mutual TLS (mTLS).", "APPROVED"),
        ("SEC-03", "Data Retention", "What is the policy for data retention and deletion upon contract exit?", "Customer operational data is retained for 90 days following contract termination, after which a certified cryptographic wipe conforming to NIST SP 800-88 Rev 1 is executed, accompanied by an auditable certificate of destruction.", "APPROVED"),
        ("SEC-04", "Compliance", "List all third-party security audits and certifications held by your organization.", "ApexFlow maintains an annual SOC 2 Type II audit covering Security, Availability, and Confidentiality conducted by Ernst & Young (issued November 2025). We are also certified under ISO/IEC 27001:2022 and ISO/IEC 27701:2019.", "APPROVED"),
        ("SEC-05", "SLA & Resilience", "What SLA and disaster recovery parameters are provided?", "We guarantee a 99.95% monthly application uptime SLA backed by financial credits. In the event of a regional disaster, our Recovery Time Objective (RTO) is under 1 hour and our Recovery Point Objective (RPO) is under 15 minutes.", "APPROVED"),
        ("SEC-06", "Rate Limits & APIs", "What are your API integration capabilities and throughput limits?", "ApexFlow provides RESTful (OpenAPI 3.1) and GraphQL APIs with an enterprise rate limit of 10,000 requests per minute per tenant. Real-time webhooks with HMAC-SHA256 signatures are supported for Slack, Jira, and ServiceNow.", "APPROVED"),
    ]

    for q in past_qa:
        ws.append(q)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx in (1, 5) else "left", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws.row_dimensions[1].height = 28

    for row_idx in range(2, len(past_qa) + 2):
        ws.row_dimensions[row_idx].height = 38
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = bold_font if col_idx in (1, 2, 5) else data_font
            cell.border = cell_border
            if is_even:
                cell.fill = alt_fill
            if col_idx in (1, 5):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (3, 4):
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 44
    ws.column_dimensions['D'].width = 56
    ws.column_dimensions['E'].width = 18

    past_rfp_path = os.path.join(OUT_DIR, "Past_Completed_RFP_ApexFlow_2025.xlsx")
    wb.save(past_rfp_path)
    print(f"Created: {past_rfp_path}")

if __name__ == "__main__":
    generate_pdf()
    generate_text_and_md()
    generate_blank_rfp_excel()
    generate_past_rfp_excel()
    print("All sample test documents generated successfully!")
